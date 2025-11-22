import re
from typing import List
import io
from docx import Document

def chunk_text(text: str, chunk_size: int = 800, overlap: int = 100) -> List[str]:
    """
    Split text into overlapping chunks to maintain context.
    
    Args:
        text: Input text to chunk
        chunk_size: Target size of each chunk in characters
        overlap: Number of overlapping characters between chunks
    
    Returns:
        List of text chunks
    """
    if not text:
        return []
    
    chunks = []
    start = 0
    text_length = len(text)
    
    while start < text_length:
        end = start + chunk_size
        
        # Try to break at sentence boundary if possible
        if end < text_length:
            # Look for sentence endings near the chunk boundary
            search_start = max(start, end - 100)
            search_text = text[search_start:end + 100]
            
            # Find last sentence ending
            sentence_endings = [m.end() for m in re.finditer(r'[.!?]\s+', search_text)]
            if sentence_endings:
                last_ending = sentence_endings[-1]
                end = search_start + last_ending
        
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        
        # Move to next chunk with overlap
        start = end - overlap
        if start <= 0:
            start = end
    
    return chunks

def parse_txt(file_bytes: bytes) -> str:
    """Parse .txt file with support for multiple encodings"""
    encodings = ['utf-8', 'utf-8-sig', 'latin-1', 'cp1252', 'iso-8859-1']
    
    for encoding in encodings:
        try:
            return file_bytes.decode(encoding)
        except (UnicodeDecodeError, AttributeError):
            continue
    
    # Fallback: decode with errors='ignore'
    return file_bytes.decode('utf-8', errors='ignore')

def parse_pdf(file_bytes: bytes) -> str:
    """Parse .pdf file with multilingual support"""
    try:
        import pypdf
        pdf_file = io.BytesIO(file_bytes)
        pdf_reader = pypdf.PdfReader(pdf_file)
        
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text() + "\n"
        
        return text
    except Exception as e:
        raise ValueError(f"Failed to parse PDF: {str(e)}")

def parse_docx(file_bytes: bytes) -> str:
    """Parse .docx file with multilingual support"""
    try:
        docx_file = io.BytesIO(file_bytes)
        doc = Document(docx_file)
        
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += cell.text + " "
                text += "\n"
        
        return text
    except Exception as e:
        raise ValueError(f"Failed to parse DOCX: {str(e)}")

def detect_file_type(filename: str) -> str:
    """Detect file type from filename"""
    filename_lower = filename.lower()
    if filename_lower.endswith('.txt'):
        return 'txt'
    elif filename_lower.endswith('.pdf'):
        return 'pdf'
    elif filename_lower.endswith('.docx'):
        return 'docx'
    elif filename_lower.endswith('.doc'):
        return 'doc'
    else:
        return 'unknown'

def process_document(file_bytes: bytes, filename: str) -> tuple[str, List[str]]:
    """
    Process uploaded document: parse and chunk.
    Supports multilingual documents (English, French, Arabic, etc.)
    
    Returns:
        Tuple of (full_text, chunks)
    """
    file_type = detect_file_type(filename)
    
    if file_type == 'txt':
        text = parse_txt(file_bytes)
    elif file_type == 'pdf':
        text = parse_pdf(file_bytes)
    elif file_type == 'docx':
        text = parse_docx(file_bytes)
    elif file_type == 'doc':
        raise ValueError("Legacy .doc files are not supported. Please convert to .docx format.")
    else:
        raise ValueError(f"Unsupported file type. Supported formats: .txt, .pdf, .docx")
    
    if not text or not text.strip():
        raise ValueError("Document appears to be empty or could not be read.")
    
    chunks = chunk_text(text)
    return text, chunks
